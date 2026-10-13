import os
from pathlib import Path
from typing import List, Any


class DocumentLoader:
    SUPPORTED_PDF = ['.pdf']
    SUPPORTED_HTML = ['.html', '.htm']
    SUPPORTED_TEXT = ['.txt', '.md']
    SUPPORTED_DATA = ['.csv', '.json', '.xml']
    SUPPORTED_OFFICE = ['.docx', '.xlsx', '.pptx']

    def _document(self, text: str, metadata: dict):
        from langchain.schema import Document
        return Document(page_content=text, metadata=metadata)

    def load_pdf(self, file_path: str) -> List[Any]:
        from langchain_community.document_loaders import PyPDFLoader
        return PyPDFLoader(file_path).load()

    def _load_text(self, file_path: str) -> List[Any]:
        text = Path(file_path).read_text(encoding='utf-8', errors='replace')
        return [self._document(text, {'source': str(file_path), 'page': 0})]

    def _load_html(self, file_path: str) -> List[Any]:
        from bs4 import BeautifulSoup
        raw = Path(file_path).read_text(encoding='utf-8', errors='replace')
        soup = BeautifulSoup(raw, 'html.parser')
        for tag in soup(['script', 'style', 'nav', 'footer', 'header', 'aside']):
            tag.decompose()
        text = soup.get_text(separator='\n', strip=True)
        return [self._document(text, {'source': str(file_path), 'page': 0})]

    def _load_office(self, file_path: str) -> List[Any]:
        path, ext = Path(file_path), Path(file_path).suffix.lower()
        if ext == '.docx':
            from docx import Document
            text = '\n'.join(p.text for p in Document(path).paragraphs if p.text.strip())
        elif ext == '.xlsx':
            from openpyxl import load_workbook
            book = load_workbook(path, read_only=True, data_only=True)
            text = '\n'.join(' | '.join(str(v) for v in row if v is not None)
                             for sheet in book.worksheets for row in sheet.iter_rows(values_only=True))
        else:
            from pptx import Presentation
            text = '\n'.join(shape.text for slide in Presentation(path).slides
                             for shape in slide.shapes if hasattr(shape, 'text') and shape.text.strip())
        return [self._document(text, {'source': str(path), 'page': 0})]

    def load_directory(self, dir_path: str, glob_pattern: str = '**/*') -> List[Any]:
        root = Path(dir_path)
        if not root.exists():
            return []
        docs: List[Any] = []
        for path in sorted(root.rglob('*')):
            if not path.is_file():
                continue
            # Internal state and the derived graph are not institute source records.
            if any(part.startswith('.') for part in path.relative_to(root).parts):
                continue
            if path.name == 'knowledge_graph.json':
                continue
            try:
                docs.extend(self.load_single_file(str(path)))
            except Exception as exc:
                print(f'Skipping {path}: {exc}')
        return docs

    def load_single_file(self, file_path: str) -> List[Any]:
        ext = os.path.splitext(file_path)[1].lower()
        if ext in self.SUPPORTED_PDF:
            return self.load_pdf(file_path)
        if ext in self.SUPPORTED_HTML:
            return self._load_html(file_path)
        if ext in self.SUPPORTED_TEXT:
            return self._load_text(file_path)
        if ext in self.SUPPORTED_DATA:
            return self._load_text(file_path)
        if ext in self.SUPPORTED_OFFICE:
            return self._load_office(file_path)
        raise ValueError(f'Unsupported file format: {ext}')
